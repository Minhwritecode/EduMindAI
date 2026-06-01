from fastapi import APIRouter, HTTPException, File, UploadFile, Form
from pydantic import BaseModel, Field
from datetime import datetime, timezone
import uuid
import os
from typing import List, Optional, Dict, Any
import json
import google.generativeai as genai
import io
import pypdf
import docx
import pptx

import server.database as database

router = APIRouter(prefix="/api/notebooks", tags=["notebooks"])

# Configure Gemini
api_key = os.getenv("GEMINI_API_KEY", "").strip()
if api_key:
    genai.configure(api_key=api_key)

# Pydantic Schemas matching database spec
class SourceItem(BaseModel):
    source_id: str
    title: str
    source_type: str = "text"
    content: str = ""
    file_url: Optional[str] = ""

class MindmapNodeData(BaseModel):
    label: str

class MindmapNodePosition(BaseModel):
    x: float
    y: float

class MindmapNode(BaseModel):
    id: str
    type: Optional[str] = "default"
    data: MindmapNodeData
    position: MindmapNodePosition

class MindmapEdge(BaseModel):
    id: str
    source: str
    target: str

class MindmapGraphData(BaseModel):
    nodes: List[MindmapNode] = []
    edges: List[MindmapEdge] = []

class MindmapItem(BaseModel):
    mindmap_id: str
    title: str
    graph_data: MindmapGraphData

class QuizQuestion(BaseModel):
    question_text: str
    options: List[str]
    correct_answer: str

class QuizItem(BaseModel):
    quiz_id: str
    title: str
    questions: List[QuizQuestion]

class FlashcardItem(BaseModel):
    front: str
    back: str

class FlashcardDeck(BaseModel):
    deck_id: str
    title: str
    cards: List[FlashcardItem]

# Core Notebook Request/Response schemas
class NotebookBase(BaseModel):
    id: Optional[str] = ""
    title: str = "Untitled Notebook"
    description: str = ""
    text: str = ""  # Legacy text for backward compatibility with sources fallback
    sources: List[SourceItem] = []
    mindmaps: List[MindmapItem] = []
    quizzes: List[QuizItem] = []
    flashcards: List[FlashcardDeck] = []

class NotebookRequest(NotebookBase):
    userId: str = "local"

class NotebookResponse(NotebookBase):
    userId: str

@router.get("", response_model=dict)
async def get_notebooks(userId: str = "local"):
    if database.db is None:
        raise HTTPException(status_code=503, detail="MongoDB not configured")
    
    cursor = database.db["notebooks"].find({"userId": userId}, {"_id": 0})
    docs = await cursor.to_list(length=100)
    
    # Ensure all new array fields exist in returned docs
    formatted = []
    for doc in docs:
        if "sources" not in doc:
            # Fallback legacy text content to a source
            doc["sources"] = []
            if doc.get("text"):
                doc["sources"].append({
                    "source_id": str(uuid.uuid4()),
                    "title": "Legacy Content.txt",
                    "source_type": "text",
                    "content": doc.get("text"),
                    "file_url": ""
                })
        doc.setdefault("mindmaps", [])
        doc.setdefault("quizzes", [])
        doc.setdefault("flashcards", [])
        doc.setdefault("description", "")
        formatted.append(doc)
        
    return {"ok": True, "notebooks": formatted}

@router.post("", response_model=dict)
async def save_notebook(notebook: NotebookRequest):
    if database.db is None:
        raise HTTPException(status_code=503, detail="MongoDB not configured")
    
    notebook_id = notebook.id
    if not notebook_id:
        notebook_id = str(uuid.uuid4())
        
    # Get existing notebook if any to preserve fields not sent
    existing = await database.db["notebooks"].find_one({"userId": notebook.userId, "id": notebook_id})
    
    update_data = {
        "userId": notebook.userId,
        "id": notebook_id,
        "title": notebook.title,
        "description": notebook.description or (existing.get("description", "") if existing else ""),
        "text": notebook.text,
        "sources": [s.model_dump() for s in notebook.sources] if notebook.sources else (existing.get("sources", []) if existing else []),
        "mindmaps": [m.model_dump() for m in notebook.mindmaps] if notebook.mindmaps else (existing.get("mindmaps", []) if existing else []),
        "quizzes": [q.model_dump() for q in notebook.quizzes] if notebook.quizzes else (existing.get("quizzes", []) if existing else []),
        "flashcards": [f.model_dump() for f in notebook.flashcards] if notebook.flashcards else (existing.get("flashcards", []) if existing else []),
        "updatedAt": datetime.now(timezone.utc),
    }

    # If first time, set createdAt
    if not existing:
        update_data["createdAt"] = datetime.now(timezone.utc)

    await database.db["notebooks"].update_one(
        {"userId": notebook.userId, "id": notebook_id},
        {"$set": update_data},
        upsert=True,
    )
    
    return {"ok": True, "id": notebook_id}

@router.delete("", response_model=dict)
async def delete_notebook(userId: str, id: str):
    if database.db is None:
        raise HTTPException(status_code=503, detail="MongoDB not configured")
    
    result = await database.db["notebooks"].delete_one({"userId": userId, "id": id})
    if result.deleted_count == 0:
        raise HTTPException(status_code=404, detail="Notebook not found")
        
    # Also clean up associated chat history
    await database.db["chat_histories"].delete_one({"notebook_id": id})
        
    return {"ok": True}

@router.post("/{notebook_id}/sources", response_model=dict)
async def add_source(notebook_id: str, userId: str, title: str, content: str, source_type: str = "text", file_url: str = ""):
    if database.db is None:
        raise HTTPException(status_code=503, detail="MongoDB not configured")
        
    notebook = await database.db["notebooks"].find_one({"userId": userId, "id": notebook_id})
    if not notebook:
        raise HTTPException(status_code=404, detail="Notebook not found")
        
    source_id = str(uuid.uuid4())
    new_source = {
        "source_id": source_id,
        "title": title,
        "source_type": source_type,
        "content": content,
        "file_url": file_url
    }
    
    await database.db["notebooks"].update_one(
        {"userId": userId, "id": notebook_id},
        {"$push": {"sources": new_source}, "$set": {"updatedAt": datetime.now(timezone.utc)}}
    )
    
    return {"ok": True, "source_id": source_id}

# LLM Generation Endpoints using response_mime_type for exact structures
@router.post("/{notebook_id}/generate/mindmap", response_model=dict)
async def generate_mindmap(notebook_id: str, userId: str, title: str = "Database Architecture Map"):
    if not api_key:
        raise HTTPException(status_code=500, detail="Gemini API Key not configured")
    if database.db is None:
        raise HTTPException(status_code=503, detail="MongoDB not configured")
        
    notebook = await database.db["notebooks"].find_one({"userId": userId, "id": notebook_id})
    if not notebook:
        raise HTTPException(status_code=404, detail="Notebook not found")
        
    # Gather source texts
    sources = notebook.get("sources", [])
    context_text = "\n\n".join([s.get("content", "") for s in sources])
    if not context_text:
        context_text = notebook.get("text", "")
        
    if not context_text.strip():
        raise HTTPException(status_code=400, detail="No source content found in this notebook. Add some sources first!")

    prompt = f"""
    You are an expert educational content creator. Based on the following study materials, generate a structured hierarchal mindmap of key concepts in Vietnamese.
    
    Return a valid JSON object matching this structure EXACTLY:
    {{
      "nodes": [
        {{ "id": "1", "type": "input", "data": {{ "label": "Main Topic" }}, "position": {{ "x": 250, "y": 5 }} }},
        {{ "id": "2", "type": "default", "data": {{ "label": "Subtopic A" }}, "position": {{ "x": 100, "y": 100 }} }},
        {{ "id": "3", "type": "default", "data": {{ "label": "Subtopic B" }}, "position": {{ "x": 400, "y": 100 }} }}
      ],
      "edges": [
        {{ "id": "e1-2", "source": "1", "target": "2" }},
        {{ "id": "e1-3", "source": "1", "target": "3" }}
      ]
    }}
    
    Make sure nodes are spaced logically horizontally (x: 50 to 500) and vertically (y: 5 to 400) so they don't overlap in the viewer.
    Create between 5 and 10 nodes capturing key concepts and sub-concepts.
    
    STUDY MATERIALS:
    {context_text}
    """
    
    import time
    try:
        model = genai.GenerativeModel("gemini-flash-latest")
        
        # Implement a simple retry for rate limits
        for attempt in range(3):
            try:
                response = model.generate_content(
                    prompt,
                    generation_config={"response_mime_type": "application/json"}
                )
                break
            except Exception as e:
                if "429" in str(e) and attempt < 2:
                    time.sleep(10)
                else:
                    raise e
        
        graph_data = json.loads(response.text)
        
        mindmap_id = str(uuid.uuid4())
        new_mindmap = {
            "mindmap_id": mindmap_id,
            "title": title,
            "graph_data": graph_data
        }
        
        await database.db["notebooks"].update_one(
            {"userId": userId, "id": notebook_id},
            {"$push": {"mindmaps": new_mindmap}, "$set": {"updatedAt": datetime.now(timezone.utc)}}
        )
        return {"ok": True, "mindmap": new_mindmap}
    except Exception as e:
        raise HTTPException(status_code=500, detail=f"LLM Generation failed: {str(e)}")

@router.post("/{notebook_id}/generate/quiz", response_model=dict)
async def generate_quiz(notebook_id: str, userId: str, title: str = "Quiz trắc nghiệm"):
    if not api_key:
        raise HTTPException(status_code=500, detail="Gemini API Key not configured")
    if database.db is None:
        raise HTTPException(status_code=503, detail="MongoDB not configured")
        
    notebook = await database.db["notebooks"].find_one({"userId": userId, "id": notebook_id})
    if not notebook:
        raise HTTPException(status_code=404, detail="Notebook not found")
        
    sources = notebook.get("sources", [])
    context_text = "\n\n".join([s.get("content", "") for s in sources])
    if not context_text:
        context_text = notebook.get("text", "")
        
    if not context_text.strip():
        raise HTTPException(status_code=400, detail="No source content found in this notebook.")

    prompt = f"""
    You are an expert teacher. Based on the study material, generate an active learning quiz containing 5 challenging multiple-choice questions in Vietnamese.
    
    Return a valid JSON list of questions matching this schema:
    [
      {{
        "question_text": "What is ...?",
        "options": ["Option A", "Option B", "Option C", "Option D"],
        "correct_answer": "Option A"
      }}
    ]
    
    Make sure options are distinct and the correct_answer matches exactly one of the values inside options.
    
    STUDY MATERIALS:
    {context_text}
    """
    
    import time
    try:
        model = genai.GenerativeModel("gemini-flash-latest")
        
        for attempt in range(3):
            try:
                response = model.generate_content(
                    prompt,
                    generation_config={"response_mime_type": "application/json"}
                )
                break
            except Exception as e:
                if "429" in str(e) and attempt < 2:
                    time.sleep(10)
                else:
                    raise e
                    
        questions = json.loads(response.text)
        
        quiz_id = str(uuid.uuid4())
        new_quiz = {
            "quiz_id": quiz_id,
            "title": title,
            "questions": questions
        }
        
        await database.db["notebooks"].update_one(
            {"userId": userId, "id": notebook_id},
            {"$push": {"quizzes": new_quiz}, "$set": {"updatedAt": datetime.now(timezone.utc)}}
        )
        return {"ok": True, "quiz": new_quiz}
    except Exception as e:
        raise HTTPException(status_code=500, detail=f"LLM Generation failed: {str(e)}")

@router.post("/{notebook_id}/generate/flashcards", response_model=dict)
async def generate_flashcards(notebook_id: str, userId: str, title: str = "Flashcards"):
    if not api_key:
        raise HTTPException(status_code=500, detail="Gemini API Key not configured")
    if database.db is None:
        raise HTTPException(status_code=503, detail="MongoDB not configured")
        
    notebook = await database.db["notebooks"].find_one({"userId": userId, "id": notebook_id})
    if not notebook:
        raise HTTPException(status_code=404, detail="Notebook not found")
        
    sources = notebook.get("sources", [])
    context_text = "\n\n".join([s.get("content", "") for s in sources])
    if not context_text:
        context_text = notebook.get("text", "")
        
    if not context_text.strip():
        raise HTTPException(status_code=400, detail="No source content found in this notebook.")

    prompt = f"""
    Based on the study materials, extract key terminology, facts, and definitions to create a set of 8 useful flashcards in Vietnamese.
    
    Return a valid JSON list of cards matching this schema:
    [
      {{
        "front": "Term or question",
        "back": "Detailed definition or answer"
      }}
    ]
    
    Make the front prompt concise, and the back side clear and explanatory.
    
    STUDY MATERIALS:
    {context_text}
    """
    
    import time
    try:
        model = genai.GenerativeModel("gemini-flash-latest")
        
        for attempt in range(3):
            try:
                response = model.generate_content(
                    prompt,
                    generation_config={"response_mime_type": "application/json"}
                )
                break
            except Exception as e:
                if "429" in str(e) and attempt < 2:
                    time.sleep(10)
                else:
                    raise e
                    
        cards = json.loads(response.text)
        
        deck_id = str(uuid.uuid4())
        new_deck = {
            "deck_id": deck_id,
            "title": title,
            "cards": cards
        }
        
        await database.db["notebooks"].update_one(
            {"userId": userId, "id": notebook_id},
            {"$push": {"flashcards": new_deck}, "$set": {"updatedAt": datetime.now(timezone.utc)}}
        )
        return {"ok": True, "flashcard": new_deck}
    except Exception as e:
        raise HTTPException(status_code=500, detail=f"LLM Generation failed: {str(e)}")


def extract_text_from_bytes(file_bytes: bytes, filename: str) -> str:
    ext = filename.split(".")[-1].lower()
    
    if ext == "pdf":
        reader = pypdf.PdfReader(io.BytesIO(file_bytes))
        text_list = []
        for page in reader.pages:
            t = page.extract_text()
            if t:
                text_list.append(t)
        return "\n\n".join(text_list)
        
    elif ext == "docx":
        doc = docx.Document(io.BytesIO(file_bytes))
        text_list = []
        for paragraph in doc.paragraphs:
            if paragraph.text:
                text_list.append(paragraph.text)
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    if cell.text:
                        text_list.append(cell.text)
        return "\n\n".join(text_list)
        
    elif ext == "pptx":
        prs = pptx.Presentation(io.BytesIO(file_bytes))
        text_list = []
        for slide in prs.slides:
            for shape in slide.shapes:
                if hasattr(shape, "text") and shape.text:
                    text_list.append(shape.text)
        return "\n\n".join(text_list)
        
    else:
        # Fallback to general plain text decoding
        try:
            return file_bytes.decode("utf-8")
        except Exception:
            try:
                return file_bytes.decode("latin-1")
            except Exception as e:
                raise ValueError(f"Unable to decode text file: {str(e)}")


@router.post("/{notebook_id}/upload", response_model=dict)
async def upload_document(notebook_id: str, userId: str = Form(...), file: UploadFile = File(...)):
    if database.db is None:
        raise HTTPException(status_code=503, detail="MongoDB not configured")
        
    notebook = await database.db["notebooks"].find_one({"userId": userId, "id": notebook_id})
    if not notebook:
        raise HTTPException(status_code=404, detail="Notebook not found")
        
    try:
        file_bytes = await file.read()
        extracted_text = extract_text_from_bytes(file_bytes, file.filename)
        
        if not extracted_text.strip():
            raise HTTPException(status_code=400, detail="Document contains no readable text content.")
            
        source_id = str(uuid.uuid4())
        new_source = {
            "source_id": source_id,
            "title": file.filename,
            "source_type": file.filename.split(".")[-1].lower(),
            "content": extracted_text,
            "file_url": ""
        }
        
        await database.db["notebooks"].update_one(
            {"userId": userId, "id": notebook_id},
            {"$push": {"sources": new_source}, "$set": {"updatedAt": datetime.now(timezone.utc)}}
        )
        
        return {"ok": True, "source_id": source_id, "title": file.filename}
    except ValueError as ve:
        raise HTTPException(status_code=400, detail=str(ve))
    except Exception as e:
        raise HTTPException(status_code=500, detail=f"Failed to parse uploaded document: {str(e)}")
